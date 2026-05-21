"""Build a PPTX deck embedding the rendered slide PNG full-bleed."""
from pptx import Presentation
from pptx.util import Inches, Emu

PNG = r"c:\MAS_final\rca-mas\docs\three-way-comparison-slide.png"
PPTX = r"c:\MAS_final\rca-mas\docs\three-way-comparison-slide.pptx"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(6.667)

blank_layout = prs.slide_layouts[6]
slide = prs.slides.add_slide(blank_layout)
slide.shapes.add_picture(PNG, 0, 0, width=prs.slide_width, height=prs.slide_height)

prs.save(PPTX)
print(f"Wrote {PPTX}")
